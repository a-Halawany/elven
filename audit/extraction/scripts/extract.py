import sys, os, json, glob
from pypdf import PdfReader
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
S = sys.argv[1]
stats = {}

def iter_block_items(doc):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, doc)
        elif child.tag == qn('w:tbl'):
            yield Table(child, doc)

for src in sorted(glob.glob(os.path.join(S,"src","v*.*"))):
    tag, ext = os.path.basename(src).rsplit(".",1)
    out = os.path.join(S, f"{tag}.txt")
    if ext == "pdf":
        r = PdfReader(src)
        pages = []
        with open(out,"w",encoding="utf-8") as f:
            for i, p in enumerate(r.pages, 1):
                try:
                    t = p.extract_text() or ""
                except Exception as e:
                    t = f"[[EXTRACTION ERROR: {e!r}]]"
                pages.append(len(t.strip()))
                f.write(f"=== PAGE {i} ===\n{t}\n")
        stats[tag] = {"kind":"pdf","units":len(r.pages),"unit_chars":pages,
                      "chars":os.path.getsize(out),
                      "meta":{k:str(v) for k,v in (r.metadata or {}).items()}}
        print(tag, "pdf pages", len(r.pages), "encrypted", r.is_encrypted)
    else:
        d = docx.Document(src)
        n = 0
        rows_written = 0
        with open(out,"w",encoding="utf-8") as f:
            for blk in iter_block_items(d):
                if isinstance(blk, Paragraph):
                    n += 1
                    if n % 50 == 1:
                        f.write(f"=== PARA {n} ===\n")
                    sty = (blk.style.name if blk.style is not None else "") or ""
                    txt = blk.text
                    if sty.lower().startswith("heading") or sty.lower()=="title":
                        f.write("## " + txt + "\n")
                    else:
                        f.write(txt + "\n")
                else:
                    f.write("[TABLE]\n")
                    for row in blk.rows:
                        cells = []
                        seen = set()
                        for c in row.cells:
                            if id(c._tc) in seen: continue
                            seen.add(id(c._tc))
                            cells.append(c.text.replace("\n"," / "))
                        f.write(" | ".join(cells) + "\n")
                        rows_written += 1
                    f.write("[/TABLE]\n")
        stats[tag] = {"kind":"docx","units":n,"tables":len(d.tables),"table_rows":rows_written,
                      "chars":os.path.getsize(out)}
        print(tag, "docx paras", n, "tables", len(d.tables), "rows", rows_written)
json.dump(stats, open(os.path.join(S,"stats.json"),"w"), indent=1)
